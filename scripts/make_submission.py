"""Папка для сдачи в форму: презентация (.pptx + .pdf), демо-видео, документ со сценариями.

Слайды для .pptx — картинки 1920×1080 из docs/slides/okno_slides.html (тот же дизайн, что в PDF), текст выступления — в заметках к слайдам.
Запуск: python scripts/make_submission.py  →  submission/
"""
import re
import shutil
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "submission"
SLIDES = ROOT / "docs" / "slides"


def notes_from_pitch() -> list[str]:
    """Текст выступления по слайдам из docs/pitch.md (строки «**N. Заголовок …**» и абзац после)."""
    t = (ROOT / "docs" / "pitch.md").read_text(encoding="utf-8")
    parts = re.split(r"\n\*\*(\d+)\. ", t)
    notes = {}
    for i in range(1, len(parts) - 1, 2):
        notes[int(parts[i])] = re.split(r"\n---|\n### ", parts[i + 1])[0].strip()
    return [notes.get(i, "") for i in range(1, 13)]


def main():
    OUT.mkdir(exist_ok=True)
    imgs = []
    with sync_playwright() as p:
        b = p.chromium.launch(channel="chrome", headless=True)
        pg = b.new_page(viewport={"width": 1920, "height": 1080})
        pg.goto((SLIDES / "okno_slides.html").as_uri())
        pg.wait_for_timeout(2500)
        for i, el in enumerate(pg.query_selector_all("section.s"), 1):
            f = OUT / f"_slide{i:02d}.jpg"
            el.screenshot(path=str(f), type="jpeg", quality=88)
            imgs.append(f)
        b.close()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9
    notes = notes_from_pitch()
    for i, f in enumerate(imgs):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(f), 0, 0, width=prs.slide_width, height=prs.slide_height)
        if i < len(notes) and notes[i]:
            s.notes_slide.notes_text_frame.text = notes[i]
    prs.save(OUT / "Okno_presentation.pptx")
    for f in imgs:
        f.unlink()
    shutil.copy(ROOT / "docs" / "Okno_presentation.pdf", OUT / "Okno_presentation.pdf")
    shutil.copy(ROOT / "docs" / "Okno_scenarios.pdf", OUT / "Okno_test_scenarios.pdf")
    if (ROOT / "video" / "okno_demo.mp4").exists():
        shutil.copy(ROOT / "video" / "okno_demo.mp4", OUT / "Okno_demo.mp4")
    for f in sorted(OUT.iterdir()):
        print(f"{f.name:28} {f.stat().st_size / 1e6:6.1f} MB")


if __name__ == "__main__":
    main()
